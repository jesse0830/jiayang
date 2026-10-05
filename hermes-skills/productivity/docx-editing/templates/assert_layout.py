# 交付"优化版 docx"时的交付前断言模板
# 用途：本机没有 LibreOffice、WPS headless 转 PDF 会挂住时，用回读断言代替目视检查。
# 用法：改 OUT 路径与预期值后 PYTHONPATH= python3 assert_layout.py
# 只断言"结构性事实"（样式/字体/对齐/缩进/表格数/图形段行数与字体），不做"好不好看"判断。
from docx import Document
from docx.shared import Pt

OUT = '/path/to/优化版.docx'
d = Document(OUT)
paras = d.paragraphs


def pstyle(i):
    p = paras[i]
    return (p.style.name, (p.text or '')[:30], p.alignment,
            p.paragraph_format.first_line_indent)


# 1) 封面 / 目录 / 导览页应为各自的样式段（数量按实际填）
print('cover:', [pstyle(i) for i in range(0, 8)])

# 2) 标题层级：一级=章(黑体16/14/12pt)、二级=x.y、三级
for i in range(0, min(120, len(paras))):
    p = paras[i]
    if p.style.name.startswith('Heading'):
        print(i, p.style.name, p.text[:40])

# 3) 表格数：真数据表应保留，1×1 强调框应归零
t1 = [t for t in d.tables if len(t.rows) == 1 and len(t.columns) == 1]
print('tables:', len(d.tables), '| 1x1 boxes left:', len(t1))  # 期望 boxes=0

# 4) 保留的 ASCII 分叉图：行数 + 是否等宽字体
for i, p in enumerate(paras):
    if p.text.rstrip().endswith('┘') or '┌' in p.text:
        fonts = {r.font.name for r in p.runs}
        print('forkgraph', i, 'lines=', p.text.count('\n') + 1, 'fonts=', fonts)
        # 期望 fonts 含 'Consolas'（中文部分 eastAsia 另设）

# 5) 内容缺失标记位（红字提示）数量应与说明文件一致
marks = [i for i, p in enumerate(paras) if '原文档此处' in p.text]
print('missing-content marks:', len(marks), marks)

# 6) 残留清理：不应再出现导出残留标记
bad = [i for i, p in enumerate(paras) if p.text.strip() == 'Plain Text']
print('residual "Plain Text":', len(bad))  # 期望 0
