#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""干部培训作业 docx 生成脚本模板（作业一/作业二/后续作业通用）
实测作业一 docx（用户已交版本）格式：
  标题(黑体16pt 加粗居中) / 副标题行(宋体12pt 居中)
  一级标题"一、"（黑体14pt 加粗）/ 二级标题"（一）"（宋体12pt 加粗）
  正文(宋体12pt, 西文 Times New Roman, 首行缩进 0.85cm = 304800 EMU)
  页边距: 上下 2.54cm, 左右 3.17cm
用法: 改 OUT 路径与内容, 直接运行 `python3 gen_hw.py`。
"""
from docx import Document
from docx.shared import Pt, Cm, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

OUT = "/path/to/作业N+姓名+题目_A和B都有.docx"

doc = Document()
sec = doc.sections[0]
sec.top_margin = Cm(2.54)
sec.bottom_margin = Cm(2.54)
sec.left_margin = Cm(3.17)
sec.right_margin = Cm(3.17)

def set_font(run, cn_font, size, bold=False, ascii_font="Times New Roman"):
    run.font.name = ascii_font
    run.font.size = Pt(size)
    run.font.bold = bold
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn('w:ascii'), ascii_font)
    rFonts.set(qn('w:hAnsi'), ascii_font)
    rFonts.set(qn('w:eastAsia'), cn_font)

def title(text, after=12):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(after)
    set_font(p.add_run(text), "黑体", 16, True)

def subtitle(text, after=18):
    """副标题/署名行, 如 '——…培训心得与思考\\n杨嘉阳 2026年8月'"""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(after)
    set_font(p.add_run(text), "宋体", 12, False)

def h1(text, before=14, after=8):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    set_font(p.add_run(text), "黑体", 14, True)

def h2(text, before=8, after=4):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    set_font(p.add_run(text), "宋体", 12, True)

def body(text, indent=True):
    p = doc.add_paragraph()
    if indent:
        p.paragraph_format.first_line_indent = Emu(304800)
    set_font(p.add_run(text), "宋体", 12, False)

# ---------------- 内容从这里写 ----------------
title("作业二：组织管理机制构建作业")
subtitle("——上下同欲：战略澄清与目标管理 培训心得与思考\n杨嘉阳    2026年8月")
body("引言段……")
h1("一、……（A题）")
h2("（一）……")
body("正文……")
# ---------------- 内容写到这里 ----------------

doc.save(OUT)
print("saved:", OUT)
