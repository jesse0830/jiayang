#!/usr/bin/env python3
"""Audit a folder of documents (PDF-focused): manifest -> text cache -> stats -> xlsx index.

Usage:
    cd /tmp && python3 corpus_audit.py <ROOT> [OUT_DIR]

ROOT     folder to audit (recursed)
OUT_DIR  where the index is written (default: ROOT)

Stages cache to /tmp/corpus_audit/<slug>/{manifest.json,text.json} and are skipped
when the cache exists, so re-running to add analysis is cheap.
"""
import os
import sys
import re
import json
import hashlib
import collections

ROOT = os.path.abspath(os.path.expanduser(sys.argv[1]))
OUT = os.path.abspath(os.path.expanduser(sys.argv[2])) if len(sys.argv) > 2 else ROOT
SLUG = re.sub(r'\W+', '_', ROOT)[-60:]
CACHE = os.path.join('/tmp/corpus_audit', SLUG)
os.makedirs(CACHE, exist_ok=True)

EXPORT_LOSS = ['\u6b64\u5904\u4e3a\u8bed\u96c0\u5185\u5bb9\u5361\u7247',
               '\u6b64\u5904\u4e3a\u8bed\u96c0',
               '\u70b9\u51fb\u94fe\u63a5\u67e5\u770b']
SENSITIVE = ['\u62a5\u4ef7', '\u4ef7\u683c', '\u5143/\u5957',
             '\u670d\u52a1\u6761\u6b3e', '\u6388\u6743\u8303\u56f4', '\u5546\u52a1']
TIMESTAMP = re.compile(r'20\d\d[/-]\d{1,2}[/-]\d{1,2}[ ]+\d{1,2}:\d{2}')
URL = re.compile(r'https?://[^\s)\uff09]+')


def cached(name, build):
    path = os.path.join(CACHE, name)
    if os.path.exists(path):
        return json.load(open(path))
    data = build()
    json.dump(data, open(path, 'w'), ensure_ascii=False)
    return data


def build_manifest():
    rows = []
    for dirpath, _dirnames, filenames in os.walk(ROOT):
        for name in sorted(filenames):
            if name.startswith('.'):
                continue
            full = os.path.join(dirpath, name)
            digest = hashlib.md5()
            with open(full, 'rb') as fh:
                for chunk in iter(lambda: fh.read(1 << 20), b''):
                    digest.update(chunk)
            rows.append({'rel': os.path.relpath(full, ROOT),
                         'size': os.path.getsize(full),
                         'md5': digest.hexdigest()})
    return rows


def pdf_rows(manifest):
    return [r for r in manifest if r['rel'].lower().endswith('.pdf')]


def extract_one(rel):
    import pymupdf
    doc = pymupdf.open(os.path.join(ROOT, rel))
    try:
        pages = doc.page_count
        text = '\n'.join(page.get_text() for page in doc)
        head = doc[0].get_text()[:180].replace('\n', ' ') if pages else ''
    finally:
        doc.close()
    return {'pages': pages, 'text': text, 'head': head, 'err': ''}


def build_text(manifest):
    rows = pdf_rows(manifest)
    out = {}
    for i, row in enumerate(rows, 1):
        try:
            out[row['rel']] = extract_one(row['rel'])
        except Exception as exc:  # one corrupt/encrypted file must not abort the scan
            out[row['rel']] = {'pages': 0, 'text': '', 'head': '',
                               'err': '%s: %s' % (type(exc).__name__, exc)}
        if i % 20 == 0:
            print('extracted %d/%d' % (i, len(rows)), flush=True)
    return out


def section_of(rel):
    parts = rel.split(os.sep)
    return parts[0] if len(parts) > 1 else '(root files)'


def collect(manifest, texts):
    rows = []
    for row in manifest:
        info = texts.get(row['rel'])
        if info is None:
            continue
        loss = sum(info['text'].count(p) for p in EXPORT_LOSS)
        links = len(URL.findall(info['text']))
        flags = ','.join(k for k in SENSITIVE if k in info['text'])
        stamps = sorted(set(TIMESTAMP.findall(info['text'])))[:3]
        rows.append((row, info, loss, links, flags, stamps))
    return rows


def write_index(rows, out_dir):
    header = ['\u5e8f\u53f7', '\u4e00\u7ea7\u7ae0\u8282', '\u4e8c\u7ea7', '\u4e09\u7ea7',
              '\u6587\u4ef6\u540d', '\u9875\u6570', '\u5927\u5c0fMB', '\u5916\u94fe\u6570',
              '\u5bfc\u51fa\u7f3a\u5931\u5904', '\u654f\u611f\u6807\u8bb0',
              '\u9996\u9875\u6458\u8981', '\u76f8\u5bf9\u8def\u5f84']

    def cells(i, row, info, loss, links, flags):
        parts = row['rel'].split(os.sep)
        return [i,
                parts[0] if len(parts) > 1 else '(root files)',
                parts[1] if len(parts) > 2 else None,
                parts[2] if len(parts) > 3 else None,
                os.path.splitext(parts[-1])[0],
                info['pages'], round(row['size'] / 1048576.0, 2),
                links, loss, flags, info['head'], row['rel']]

    try:
        from openpyxl import Workbook
        from openpyxl.utils import get_column_letter
        book = Workbook()
        sheet = book.active
        sheet.title = 'index'
        sheet.append(header)
        for i, (row, info, loss, links, flags, _stamps) in enumerate(rows, 1):
            sheet.append(cells(i, row, info, loss, links, flags))
        sheet.freeze_panes = 'A2'
        sheet.auto_filter.ref = 'A1:L%d' % sheet.max_row
        for idx, width in enumerate([5, 22, 18, 18, 34, 6, 8, 9, 11, 14, 60, 60], 1):
            sheet.column_dimensions[get_column_letter(idx)].width = width
        path = os.path.join(out_dir, 'corpus_index.xlsx')
        book.save(path)
        print('\nindex -> %s (%d rows)' % (path, sheet.max_row - 1))
    except ImportError:
        import csv
        path = os.path.join(out_dir, 'corpus_index.csv')
        with open(path, 'w', newline='', encoding='utf-8-sig') as fh:
            writer = csv.writer(fh)
            writer.writerow(header)
            for i, (row, info, loss, links, flags, _stamps) in enumerate(rows, 1):
                writer.writerow(cells(i, row, info, loss, links, flags))
        print('\nopenpyxl unavailable -> csv at %s' % path)


def main():
    manifest = cached('manifest.json', build_manifest)
    texts = cached('text.json', lambda: build_text(manifest))

    total_mb = sum(r['size'] for r in manifest) / 1048576.0
    total_pages = sum(v['pages'] for v in texts.values())
    total_chars = sum(len(v['text']) for v in texts.values())
    print('files=%d pdfs=%d pages=%d chars=%d size=%.1fMB'
          % (len(manifest), len(texts), total_pages, total_chars, total_mb))

    by_hash = collections.defaultdict(list)
    for row in manifest:
        by_hash[row['md5']].append(row['rel'])
    dups = sorted((g for g in by_hash.values() if len(g) > 1), key=len, reverse=True)
    print('\n== %d duplicate groups ==' % len(dups))
    for group in dups:
        print('  %dx %s' % (len(group), ' | '.join(group)))

    rows = collect(manifest, texts)

    agg = collections.defaultdict(lambda: [0, 0, 0, 0])
    for row, info, _loss, _links, _flags, _stamps in rows:
        slot = agg[section_of(row['rel'])]
        slot[0] += 1
        slot[1] += info['pages']
        slot[2] += row['size']
        slot[3] += len(info['text'])
    print('\n== by top-level section: files / pages / MB / chars ==')
    for key in sorted(agg):
        slot = agg[key]
        print('  %-30s %3d %6d %8.1f %9d'
              % (key, slot[0], slot[1], slot[2] / 1048576.0, slot[3]))

    print('\n== export loss / sensitive / timestamps ==')
    for row, info, loss, links, flags, stamps in rows:
        if loss or flags:
            print('  loss=%2d links=%3d %s%s%s'
                  % (loss, links, ('[' + flags + '] ') if flags else '',
                     (stamps[0] + ' ') if stamps else '', row['rel']))

    print('\n== structure gaps ==')
    numbered = collections.defaultdict(set)
    odd = []
    for dirpath, dirnames, filenames in os.walk(ROOT):
        relp = os.path.relpath(dirpath, ROOT)
        for name in list(dirnames) + list(filenames):
            match = re.match(r'^(\d{1,2})-', name)
            if match:
                numbered[relp].add(int(match.group(1)))
            elif name in dirnames:
                odd.append(os.path.join(relp, name))
    for relp in sorted(numbered):
        nums = numbered[relp]
        missing = sorted(set(range(0, max(nums) + 1)) - nums)
        if missing and relp != '.':
            print('  %s: missing %s (have %s)' % (relp, missing, sorted(nums)))
    for name in odd:
        print('  non-numbered dir: %s' % name)

    write_index(rows, OUT)


if __name__ == '__main__':
    main()
